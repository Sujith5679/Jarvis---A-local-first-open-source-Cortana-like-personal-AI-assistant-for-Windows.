from __future__ import annotations

import openpyxl
import pymupdf
from documents.base import DocumentParseError
from documents.docx import DocxParser
from documents.pdf import PdfParser
from documents.pptx import PptxParser
from documents.registry import get_parser, is_supported
from documents.text import HtmlParser, PlainTextParser
from documents.xlsx import XlsxParser

import docx
from pptx import Presentation


def test_txt_parser_extracts_content(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text("Hello from a plain text file.", encoding="utf-8")
    doc = PlainTextParser().extract(path)
    assert doc.source_type == "txt"
    assert "Hello from a plain text file." in doc.segments[0].text


def test_html_parser_strips_tags_and_scripts(tmp_path):
    path = tmp_path / "page.html"
    path.write_text(
        "<html><head><script>evil()</script></head>"
        "<body><h1>Title</h1><p>Real content.</p></body></html>",
        encoding="utf-8",
    )
    doc = HtmlParser().extract(path)
    text = doc.segments[0].text
    assert "Real content." in text
    assert "Title" in text
    assert "evil()" not in text


def test_pdf_parser_extracts_pages(tmp_path):
    path = tmp_path / "doc.pdf"
    pdf = pymupdf.open()
    page1 = pdf.new_page()
    page1.insert_text((72, 72), "First page content about fraud detection.")
    page2 = pdf.new_page()
    page2.insert_text((72, 72), "Second page content about evaluation results.")
    pdf.save(str(path))
    pdf.close()

    doc = PdfParser().extract(path)
    assert doc.source_type == "pdf"
    assert len(doc.segments) == 2
    assert doc.segments[0].page == 1
    assert "fraud detection" in doc.segments[0].text
    assert doc.segments[1].page == 2
    assert "evaluation results" in doc.segments[1].text


def test_pdf_parser_raises_on_corrupt_file(tmp_path):
    import pytest

    path = tmp_path / "corrupt.pdf"
    path.write_bytes(b"not a real pdf")
    with pytest.raises(DocumentParseError):
        PdfParser().extract(path)


def test_docx_parser_extracts_headings_and_body(tmp_path):
    path = tmp_path / "doc.docx"
    document = docx.Document()
    document.add_heading("Risk Assessment", level=1)
    document.add_paragraph("The risk assessment found no major issues.")
    document.save(str(path))

    doc = DocxParser().extract(path)
    assert doc.source_type == "docx"
    assert any(s.section == "Risk Assessment" for s in doc.segments)
    assert any("no major issues" in s.text for s in doc.segments)


def test_pptx_parser_extracts_slides(tmp_path):
    path = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide_layout = presentation.slide_layouts[1]
    slide = presentation.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Quarterly Results"
    presentation.save(str(path))

    doc = PptxParser().extract(path)
    assert doc.source_type == "pptx"
    assert len(doc.segments) == 1
    assert doc.segments[0].page == 1
    assert "Quarterly Results" in doc.segments[0].text


def test_xlsx_parser_extracts_sheets(tmp_path):
    path = tmp_path / "sheet.xlsx"
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Revenue"
    ws.append(["Quarter", "Revenue"])
    ws.append(["Q1", 1000])
    workbook.save(str(path))

    doc = XlsxParser().extract(path)
    assert doc.source_type == "xlsx"
    assert doc.segments[0].section == "Revenue"
    assert "Revenue" in doc.segments[0].text
    assert "1000" in doc.segments[0].text


def test_registry_dispatches_by_extension(tmp_path):
    txt_path = tmp_path / "a.txt"
    assert is_supported(txt_path) is True
    assert isinstance(get_parser(txt_path), PlainTextParser)

    unknown_path = tmp_path / "a.exe"
    assert is_supported(unknown_path) is False
    assert get_parser(unknown_path) is None
