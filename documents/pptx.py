"""PPTX extraction via python-pptx. One segment per slide (page = slide number)."""

from __future__ import annotations

from pathlib import Path

from documents.base import DocumentParseError, ExtractedDocument, ExtractedSegment
from pptx import Presentation


class PptxParser:
    extensions = (".pptx",)

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            presentation = Presentation(str(path))
        except Exception as exc:
            raise DocumentParseError(f"Could not open PPTX: {exc}") from exc

        segments: list[ExtractedSegment] = []
        for slide_index, slide in enumerate(presentation.slides):
            lines: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        lines.append(text)
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        if row_text.strip(" |"):
                            lines.append(row_text)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    lines.append(f"[Speaker notes] {notes}")

            text = "\n".join(lines).strip()
            if text:
                segments.append(ExtractedSegment(text=text, page=slide_index + 1))

        return ExtractedDocument(segments=segments, source_type="pptx")
