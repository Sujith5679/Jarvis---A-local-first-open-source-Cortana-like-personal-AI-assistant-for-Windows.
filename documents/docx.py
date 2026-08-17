"""DOCX extraction via python-docx.

Paragraphs are grouped into segments by heading ("Heading 1/2/3" styles) so
each chunk can carry a meaningful `section` label; tables are extracted as
their own segments.
"""

from __future__ import annotations

from pathlib import Path

import docx
from documents.base import DocumentParseError, ExtractedDocument, ExtractedSegment


class DocxParser:
    extensions = (".docx",)

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            document = docx.Document(str(path))
        except Exception as exc:
            raise DocumentParseError(f"Could not open DOCX: {exc}") from exc

        segments: list[ExtractedSegment] = []
        current_section: str | None = None
        current_lines: list[str] = []

        def flush() -> None:
            text = "\n".join(current_lines).strip()
            if text:
                segments.append(ExtractedSegment(text=text, section=current_section))

        for para in document.paragraphs:
            style_name = (para.style.name if para.style else "") or ""
            text = para.text.strip()
            if not text:
                continue
            if style_name.startswith("Heading"):
                flush()
                current_lines = []
                current_section = text
            else:
                current_lines.append(text)
        flush()

        for table_index, table in enumerate(document.tables):
            rows = [
                " | ".join(cell.text.strip() for cell in row.cells) for row in table.rows
            ]
            table_text = "\n".join(r for r in rows if r).strip()
            if table_text:
                segments.append(
                    ExtractedSegment(text=table_text, section=f"Table {table_index + 1}")
                )

        return ExtractedDocument(segments=segments, source_type="docx")
