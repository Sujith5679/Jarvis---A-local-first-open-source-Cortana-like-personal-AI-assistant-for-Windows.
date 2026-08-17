"""XLSX extraction via openpyxl. One segment per sheet (section = sheet name)."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from documents.base import DocumentParseError, ExtractedDocument, ExtractedSegment

MAX_ROWS_PER_SHEET = 5000  # guard against pathological spreadsheets


class XlsxParser:
    extensions = (".xlsx",)

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        except Exception as exc:
            raise DocumentParseError(f"Could not open XLSX: {exc}") from exc

        segments: list[ExtractedSegment] = []
        try:
            for sheet in workbook.worksheets:
                lines: list[str] = []
                for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if row_index >= MAX_ROWS_PER_SHEET:
                        break
                    cells = ["" if v is None else str(v) for v in row]
                    if any(c.strip() for c in cells):
                        lines.append(" | ".join(cells))
                text = "\n".join(lines).strip()
                if text:
                    segments.append(ExtractedSegment(text=text, section=sheet.title))
        finally:
            workbook.close()

        return ExtractedDocument(segments=segments, source_type="xlsx")
